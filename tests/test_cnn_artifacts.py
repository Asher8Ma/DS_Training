"""Verify the latest complete CNN run's saved results and deliverables."""

import json
from pathlib import Path
import zipfile

import numpy as np
from openpyxl import load_workbook
from PIL import Image
import pytest


@pytest.fixture(scope="module")
def completed_run() -> tuple[Path, dict]:
    """Locate the newest exported run, not an incomplete training attempt."""
    output_root = Path(__file__).resolve().parents[1] / "E - nn_intro" / "part2" / "outputs"
    workbooks = sorted(output_root.glob(pattern="cnn_*/cnn_results.xlsx"))
    assert workbooks, "No exported CNN run exists"
    directory = workbooks[-1].parent
    return directory, json.loads(s=(directory / "metrics.json").read_text(encoding="utf-8"))


def test_excel_matches_notebook(*, completed_run: tuple[Path, dict]) -> None:
    """Every exported result must match the notebook's unrounded metric record."""
    directory, payload = completed_run
    workbook = load_workbook(filename=directory / "cnn_results.xlsx", read_only=True, data_only=True)
    assert workbook.sheetnames == ["Results", "Training", "Class metrics", "Sources"]
    rows = list(workbook["Results"].iter_rows(min_row=5, max_row=19, max_col=7, values_only=True))
    assert len(payload["results"]) == len(rows) == 15
    for row, record in zip(rows, payload["results"], strict=True):
        assert row[:2] == (record["model"], record["partition"])
        expected = [record[key] for key in ("samples", "loss", "accuracy", "balanced_accuracy", "macro_auc")]
        np.testing.assert_allclose(actual=row[2:], desired=expected, rtol=1e-12, atol=1e-12)
    workbook.close()


def test_saved_model_containers(*, completed_run: tuple[Path, dict]) -> None:
    """Keras model archives must be complete and include weights/configuration."""
    directory, _ = completed_run
    for name in ("cifar_dense.keras", "cifar_cnn.keras"):
        with zipfile.ZipFile(file=directory / name) as archive:
            assert archive.testzip() is None
            assert {"config.json", "metadata.json", "model.weights.h5"}.issubset(archive.namelist())
    assert (directory / "face_best.weights.h5").stat().st_size > 100_000_000


def test_prediction_and_figure_files(*, completed_run: tuple[Path, dict]) -> None:
    """Check every model's partitions, normalized probabilities, and PNG integrity."""
    directory, payload = completed_run
    prediction_files = list(directory.glob(pattern="*_predictions.npz"))
    assert len(prediction_files) == 15
    for prediction_path in prediction_files:
        with np.load(file=prediction_path, allow_pickle=False) as arrays:
            targets = arrays["targets"]
            probabilities = arrays["probabilities"]
            assert probabilities.shape[0] == targets.size
            assert np.isfinite(probabilities).all()
            assert (probabilities >= 0).all() and (probabilities <= 1).all()
            np.testing.assert_allclose(actual=probabilities.sum(axis=1), desired=1.0, atol=1e-5)
    figures = list(directory.glob(pattern="*.png"))
    assert len(figures) >= 16
    for figure_path in figures:
        with Image.open(fp=figure_path) as figure:
            assert figure.width >= 300 and figure.height >= 100
            figure.verify()
    assert payload["frozen_weights_unchanged"] is True
    training = {record["model"]: record for record in payload["training"]}
    assert training["CIFAR CNN"]["epochs"] == 40
    assert training["VGGFace neural head"]["epochs"] == 5
